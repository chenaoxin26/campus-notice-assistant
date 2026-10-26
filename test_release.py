"""Release-specific tests with invented data only; no platform collection."""
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from app import Service
from core import CampusDB, LocalAI
from excel_export import export_snapshot
from notification_source import NotificationSource

HERE=Path(__file__).resolve().parent

class ReleaseTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder=Path(self.temp.name)

    def test_fresh_install_is_empty_and_inactive(self):
        with patch('subprocess.Popen',side_effect=AssertionError('unexpected launch')):
            service=Service(self.folder/'data')
        self.assertEqual(service.config['groups'],[])
        self.assertFalse(service.source.running())
        self.assertIsNone(service.chat.day)
        self.assertEqual(service.db.snapshot()['tasks'],[])
        self.assertEqual(Path(service.config['collector_root']),HERE/'collector')
        self.assertEqual(Path(service.config['collector_data_dir']),self.folder/'data'/'collector')

    def test_empty_whitelist_cannot_start_collection(self):
        service=Service(self.folder/'data')
        with patch('subprocess.Popen',side_effect=AssertionError('unexpected launch')):
            for action in (service.source.start,service.source.scan,service.chat.start_daytime,lambda:service.chat.capture(lambda _:None)):
                with self.assertRaises(ValueError):action()

    def test_missing_model_does_not_start_or_download(self):
        ai=LocalAI(json.loads((HERE/'settings.example.json').read_text('utf-8')),self.folder)
        with patch.object(ai,'health',return_value={'available':False}),patch('subprocess.Popen',side_effect=AssertionError('unexpected launch')):
            with self.assertRaises(RuntimeError):ai.ensure()

    def test_empty_export(self):
        path=export_snapshot({'tasks':[],'notices':[]},self.folder/'empty.xlsx')
        wb=load_workbook(path)
        self.addCleanup(wb.close)
        self.assertEqual(wb.sheetnames,['校园待办','消息来源'])
        self.assertEqual(wb['校园待办'].max_row,3)

    def test_export_values_sources_and_formula_safety(self):
        n=dict(id='synthetic',source='QQ_manual',group_name='虚构群',sender=None,sent_at=None,sent_time_text=None,collected_at='2026-10-26T18:00:00+08:00',body='=HYPERLINK("https://example.invalid","虚构")')
        task=dict(title='=1+1',category='学习任务',deadline_text='待确认',location='待确认',requirements='+虚构要求',status='done',sources=['synthetic'],needs_review=True,steps=['虚构步骤'],notes='@虚构备注')
        p=export_snapshot({'tasks':[task],'notices':[n]},self.folder/'data.xlsx')
        wb=load_workbook(p,data_only=False);self.addCleanup(wb.close)
        self.assertEqual(wb['校园待办']['A4'].value,'=1+1')
        self.assertEqual(wb['校园待办']['A4'].data_type,'s')
        self.assertEqual(wb['校园待办']['F4'].value,'已完成')
        self.assertEqual(wb['消息来源']['E2'].value,'待确认')
        self.assertEqual(wb['消息来源']['G2'].data_type,'s')
        self.assertEqual(wb['消息来源']['G2'].value,n['body'])

    def test_export_omits_deleted_and_handles_limits(self):
        from excel_export import safe_text
        self.assertNotIn('\x00',safe_text('虚构\x00文本'))
        self.assertLessEqual(len(safe_text('长'*40000)),32767)
        p=export_snapshot({'tasks':[{'status':'deleted'}],'notices':[]},self.folder/'deleted.xlsx')
        wb=load_workbook(p);self.addCleanup(wb.close)
        self.assertEqual(wb['校园待办'].max_row,3)

    def test_unavailable_model_still_allows_existing_state_and_export(self):
        service=Service(self.folder/'data')
        with patch.object(service.ai,'health',return_value={'available':False}):
            self.assertEqual(service.state()['tasks'],[])
            self.assertTrue(Path(service.export()).is_file())

if __name__=='__main__':unittest.main(verbosity=2)
